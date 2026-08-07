import os
from docx import Document
from config import REPORT_FILE_PATH

def save_to_word(date, report_text):
    # যদি আগে থেকে ফাইল থাকে, সেটি ওপেন করবে; না থাকলে নতুন বানাবে
    if os.path.exists(REPORT_FILE_PATH):
        doc = Document(REPORT_FILE_PATH)
    else:
        doc = Document()
        doc.add_heading('Development Journal - BANTAI Photonic Sim', 0)
    
    # নতুন দিনের এন্ট্রি শুরু
    doc.add_paragraph('=====================').bold = True
    doc.add_heading(date, level=2)
    
    # AI-এর জেনারেট করা টেক্সট যোগ করা
    doc.add_paragraph(report_text)
    doc.add_paragraph('-------------------------')
    
    # ফাইল সেভ করা
    doc.save(REPORT_FILE_PATH)