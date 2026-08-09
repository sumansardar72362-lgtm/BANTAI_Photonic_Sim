from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Pt


REPORT_FOLDER = Path("reports")
REPORT_FILE = REPORT_FOLDER / "Development_Journal.docx"


def add_daily_report(work):

    REPORT_FOLDER.mkdir(exist_ok=True)

    if REPORT_FILE.exists():

        document = Document(REPORT_FILE)

    else:

        document = Document()

        title = document.add_heading(
            "Development Journal",
            level=0
        )

        title.alignment = 1

    today = datetime.now().strftime("%d %B %Y")

    document.add_heading(today, level=1)

    document.add_heading(
        "Today's Work",
        level=2
    )

    for item in work:

        paragraph = document.add_paragraph()

        run = paragraph.add_run(
            "• " + item["commit"]
        )

        run.bold = True
        run.font.size = Pt(11)

        if item["files"]:

            document.add_paragraph(
                "Changed Files:"
            )

            for file_name in item["files"]:

                document.add_paragraph(
                    file_name,
                    style="List Bullet"
                )

    document.add_paragraph(
        "----------------------------------------"
    )

    document.save(REPORT_FILE)

    print(
        f"Report saved: {REPORT_FILE}"
    )